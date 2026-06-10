"""检查源模板结构"""
import openpyxl

src = r"E:\rhProject\wk-train-center-service\wk-modules\wk-module-qu\src\main\resources\excel\repo_qu_tmpl.xlsx"
wb = openpyxl.load_workbook(src, data_only=False)
print("Workbook sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== Sheet: {name} (dim: {ws.dimensions}, max_row: {ws.max_row}, max_col: {ws.max_column}) ===")
    for r in range(1, min(ws.max_row + 1, 8)):
        row = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            row.append(str(v) if v is not None else "(空)")
        print(f"  R{r}: " + " | ".join(row))
