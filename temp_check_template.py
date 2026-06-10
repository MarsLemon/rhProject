"""检查试题导入模板结构。"""
import openpyxl

path = r"E:\rhProject\wk-train-center-service\wk-modules\wk-module-qu\src\main\resources\excel\repo_qu_tmpl.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
print(f"Workbook sheets: {wb.sheetnames}")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== Sheet: {name} (dim: {ws.dimensions}, max_row: {ws.max_row}, max_col: {ws.max_column}) ===")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50), values_only=False):
        row_data = []
        for cell in row:
            v = cell.value
            if v is None:
                v = ""
            row_data.append(str(v)[:30])
        print(f"  R{row[0].row}: " + " | ".join(row_data))
